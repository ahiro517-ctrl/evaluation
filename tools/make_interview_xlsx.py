#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""個人面談シート（Excel版）を生成。判定マスタ(設定)を同梱し単体で計算が成立する。"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

# ---- 制度マスタ（constants.gs と同値）----
GRADES = [
    ['B1', 'B', '1年目', 240000, 280000, 166667, 500000, 'B・A'],
    ['B2', 'B', '2年目', 240000, 280000, 500000, 1500000, 'B・A'],
    ['B3', 'B', '3年目〜', 240000, 280000, 833334, 2500000, 'B・A'],
    ['A1', 'A1', '', 300000, 349000, 1800000, 5400000, 'B・A'],
    ['A2', 'A2', '', 350000, 400000, 2100000, 6300000, 'B・A'],
    ['L1', 'L1', '', 450000, 499000, 2700000, 8100000, 'L・M'],
    ['L2', 'L2', '', 500000, 550000, 3000000, 9000000, 'L・M'],
    ['M1', 'M1', '', 600000, 649000, 3600000, 10800000, 'L・M'],
    ['M2', 'M2', '', 650000, 700000, 3900000, 11700000, 'L・M'],
]
GP_JUDGE = [
    [0, 'F', '大きく不足している'], [1.0, 'E', '大きく改善が必要'],
    [1.5, 'D', '改善が必要'], [2.0, 'C', '据え置き水準'],
    [2.5, 'B', '目標には届かないが近い水準'], [3.0, 'A', '目標を達成している'],
    [3.5, 'S', '目標を大きく超えている'],
]
VALUE_ITEMS = ['信頼', '疑問', '熱量', '創造', '実行速度', '主体性', 'チームワーク', '挑戦', '責任']
SKILL_ITEMS = ['全体管理', '予算管理', '企画', 'PR・集客', 'クリエイティブ', '事務局', '運営制作', '進行制作', '運営現場', '進行現場']
SKILL_FIELD_COUNT = 4
PAY_MATRIX_COLS = ['要改善', '標準', '高い']
PAY_MATRIX_ROWS = ['S', 'A', 'B', 'C', 'D', 'E', 'F']
PAY_MATRIX = [['A', 'S', 'S'], ['B', 'A', 'S'], ['C', 'B', 'A'], ['D', 'C', 'B'],
              ['E', 'D', 'D'], ['F', 'E', 'E'], ['F', 'F', 'F']]
RAISE_COLS = ['B・A', 'L・M']
RAISE_ROWS = ['S', 'A', 'B', 'C', 'D', 'E', 'F']
RAISE_TABLE = [[12000, 18000], [8000, 12000], [4000, 6000], [0, 0],
               [-4000, -6000], [-8000, -12000], [-12000, -18000]]
PERIOD = '2026前期'

# 設定シート参照レンジ
GRADE_TBL = "設定!$A$3:$H$11"
GPJ_LOWER = "設定!$A$15:$A$21"
GPJ_JUDGE = "設定!$B$15:$B$21"
# 等級テーブル(A:H)と重ならないようマトリクス・改定額はK列以降
MATRIX_VAL = "設定!$L$3:$N$9"
MATRIX_ROW = "設定!$K$3:$K$9"
MATRIX_COL = "設定!$L$2:$N$2"
RAISE_VAL = "設定!$L$13:$M$19"
RAISE_ROW = "設定!$K$13:$K$19"
RAISE_COL = "設定!$L$12:$M$12"

# ---- スタイル ----
HDR = Font(bold=True, color='FFFFFF')
TITLE = Font(bold=True, color='FFFFFF', size=16)
BOLD = Font(bold=True)
GREY9 = Font(color='616161', size=9)
FILL_TITLE = PatternFill('solid', fgColor='37474F')
FILL_SEC = PatternFill('solid', fgColor='455A64')
FILL_LBL = PatternFill('solid', fgColor='ECEFF1')
FILL_HEAD = PatternFill('solid', fgColor='CFD8DC')
FILL_INPUT = PatternFill('solid', fgColor='FFFDE7')
FILL_PERIOD = PatternFill('solid', fgColor='FFF9C4')
CENTER = Alignment(horizontal='center', vertical='center')
TOPWRAP = Alignment(wrap_text=True, vertical='top')
THIN = Side(style='thin', color='B0BEC5')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = openpyxl.Workbook()

# ========================= 設定シート =========================
cfg = wb.active
cfg.title = '設定'
cfg['A1'] = '■ 等級基準テーブル'; cfg['A1'].font = BOLD
for j, h in enumerate(['等級キー', '等級', '区分', '月収下限', '月収上限', '半期基準粗利(1.0倍)', '半期目標粗利(3.0倍)', '等級グループ']):
    c = cfg.cell(row=2, column=1 + j, value=h); c.font = BOLD; c.fill = FILL_HEAD
for i, row in enumerate(GRADES):
    for j, v in enumerate(row):
        cfg.cell(row=3 + i, column=1 + j, value=v)
for i in range(len(GRADES)):
    for col in (4, 5, 6, 7):
        cfg.cell(row=3 + i, column=col).number_format = '#,##0'

cfg['A13'] = '■ 粗利判定（倍率→判定）※LOOKUP用に昇順'; cfg['A13'].font = BOLD
for j, h in enumerate(['倍率下限', '判定', '状態']):
    c = cfg.cell(row=14, column=1 + j, value=h); c.font = BOLD; c.fill = FILL_HEAD
for i, row in enumerate(GP_JUDGE):
    for j, v in enumerate(row):
        cfg.cell(row=15 + i, column=1 + j, value=v)
    cfg.cell(row=15 + i, column=1).number_format = '0.0'

# 昇降給マトリクス K1:N9（等級テーブルA:Hと重ならないよう右側）
cfg['K1'] = '■ 昇降給マトリクス'; cfg['K1'].font = BOLD
cfg.cell(row=2, column=11, value='粗利判定＼バリュー').font = BOLD
cfg.cell(row=2, column=11).fill = FILL_HEAD
for j, v in enumerate(PAY_MATRIX_COLS):
    c = cfg.cell(row=2, column=12 + j, value=v); c.font = BOLD; c.fill = FILL_HEAD; c.alignment = CENTER
for i, rl in enumerate(PAY_MATRIX_ROWS):
    c = cfg.cell(row=3 + i, column=11, value=rl); c.font = BOLD; c.fill = FILL_LBL
    for j, v in enumerate(PAY_MATRIX[i]):
        cc = cfg.cell(row=3 + i, column=12 + j, value=v); cc.alignment = CENTER

# 改定額 K11:M19
cfg['K11'] = '■ 改定額（最終判定×等級グループ・円）'; cfg['K11'].font = BOLD
cfg.cell(row=12, column=11, value='最終判定＼グループ').font = BOLD
cfg.cell(row=12, column=11).fill = FILL_HEAD
for j, v in enumerate(RAISE_COLS):
    c = cfg.cell(row=12, column=12 + j, value=v); c.font = BOLD; c.fill = FILL_HEAD; c.alignment = CENTER
for i, rl in enumerate(RAISE_ROWS):
    c = cfg.cell(row=13 + i, column=11, value=rl); c.font = BOLD; c.fill = FILL_LBL
    for j, v in enumerate(RAISE_TABLE[i]):
        cc = cfg.cell(row=13 + i, column=12 + j, value=v); cc.number_format = '+#,##0;-#,##0;±0'

# 運用情報
cfg['A43'] = '■ 運用情報'; cfg['A43'].font = BOLD
cfg['A44'] = '対象期'
cfg['B44'] = PERIOD; cfg['B44'].font = BOLD; cfg['B44'].fill = FILL_PERIOD
for col, w in {'A': 18, 'B': 12, 'C': 12, 'D': 12, 'E': 12, 'F': 20, 'G': 20, 'H': 12,
               'K': 18, 'L': 10, 'M': 10, 'N': 10}.items():
    cfg.column_dimensions[col].width = w

# ========================= 面談シート =========================
ws = wb.create_sheet('面談シート')
ws.column_dimensions['A'].width = 28
for col in ('B', 'C', 'D'):
    ws.column_dimensions[col].width = 18


def label(a1, text):
    ws[a1] = text; ws[a1].font = BOLD; ws[a1].fill = FILL_LBL


def section(row, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    c = ws.cell(row=row, column=1, value=text); c.font = HDR; c.fill = FILL_SEC


def inp(a1):
    ws[a1].fill = FILL_INPUT; ws[a1].border = BORDER


def wide_input(row):
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    c = ws.cell(row=row, column=2); c.fill = FILL_INPUT; c.alignment = TOPWRAP


# タイトル
ws.merge_cells('A1:D1')
t = ws['A1']; t.value = '個人面談シート'; t.font = TITLE; t.fill = FILL_TITLE; t.alignment = CENTER
ws.row_dimensions[1].height = 28

# 基本情報
label('A2', '対象期'); ws['B2'] = "=設定!$B$44"
label('A3', '氏名'); inp('B3')
label('A4', '社員ID'); inp('B4')
label('A5', '区分'); inp('B5')
label('A6', '等級'); inp('B6'); label('C6', 'B年次'); inp('D6')
label('A7', 'ユニット'); inp('B7'); label('C7', 'ユニット長'); inp('D7')
label('A8', '現在月給'); inp('B8'); ws['B8'].number_format = '#,##0'

# 補助セル（等級キー・等級グループ）E:F は後で非表示
ws['E1'] = '等級キー'; ws['E1'].font = GREY9
ws['E2'] = '等級グループ'; ws['E2'].font = GREY9
ws['F1'] = '=IF($B$6="","",IF($B$6="B","B"&IF($D$6="",1,MIN($D$6,3)),$B$6))'
ws['F2'] = f'=IFERROR(VLOOKUP($F$1,{GRADE_TBL},8,FALSE),"")'
ws['F1'].font = GREY9; ws['F2'].font = GREY9

# 粗利評価
section(9, '■ 粗利評価')
label('A10', '半期基準粗利'); ws['B10'] = f'=IF($F$1="","",IFERROR(VLOOKUP($F$1,{GRADE_TBL},6,FALSE),""))'
label('A11', '実績半期粗利'); inp('B11')
label('A12', '粗利倍率'); ws['B12'] = '=IF(OR($B$10="",$B$11=""),"",$B$11/$B$10)'
label('A13', '粗利判定'); ws['B13'] = f'=IF($B$12="","",LOOKUP($B$12,{GPJ_LOWER},{GPJ_JUDGE}))'
ws['B10'].number_format = '#,##0'; ws['B11'].number_format = '#,##0'
ws['B12'].number_format = '0.00"倍"'
ws['A11'].comment = None

# バリュー評価
section(15, '■ バリュー評価（各3点・9項目・27点満点）')
for j, h in enumerate(['項目', '自己', 'ユニット長']):
    c = ws.cell(row=16, column=1 + j, value=h); c.font = BOLD; c.fill = FILL_HEAD
vStart = 17
for i, it in enumerate(VALUE_ITEMS):
    ws.cell(row=vStart + i, column=1, value=it)
    for col in (2, 3):
        cc = ws.cell(row=vStart + i, column=col); cc.fill = FILL_INPUT; cc.border = BORDER; cc.alignment = CENTER
vEnd = vStart + len(VALUE_ITEMS) - 1   # 25
total = vEnd + 1   # 26
conf = vEnd + 2    # 27
vj = vEnd + 3      # 28
label(f'A{total}', '合計')
ws.cell(row=total, column=2, value=f'=IF(COUNT(B{vStart}:B{vEnd})=0,"",SUM(B{vStart}:B{vEnd}))')
ws.cell(row=total, column=3, value=f'=IF(COUNT(C{vStart}:C{vEnd})=0,"",SUM(C{vStart}:C{vEnd}))')
label(f'A{conf}', '確定合計（既定＝ユニット長／会議で上書き可）')
ws.cell(row=conf, column=2, value=f'=IF($B$4="","",IF(C{total}<>"",C{total},B{total}))')
label(f'A{vj}', 'バリュー判定')
ws.cell(row=vj, column=2, value=f'=IF($B${conf}="","",IF($B${conf}>=23,"高い",IF($B${conf}>=17,"標準","要改善")))')

# スキル評価
sTitle = vj + 2  # 30
section(sTitle, '■ スキル評価（1〜5・10項目）※昇給額には使わない')
sHead = sTitle + 1  # 31
for j, h in enumerate(['項目', '自己', 'ユニット長', '確定']):
    c = ws.cell(row=sHead, column=1 + j, value=h); c.font = BOLD; c.fill = FILL_HEAD
skStart = sHead + 1  # 32
for i, it in enumerate(SKILL_ITEMS):
    r = skStart + i
    ws.cell(row=r, column=1, value=it)
    for col in (2, 3):
        cc = ws.cell(row=r, column=col); cc.fill = FILL_INPUT; cc.border = BORDER; cc.alignment = CENTER
    ws.cell(row=r, column=4, value=f'=IF($B$4="","",IF(C{r}<>"",C{r},B{r}))').alignment = CENTER
skEnd = skStart + len(SKILL_ITEMS) - 1  # 41
fldStart = skEnd - SKILL_FIELD_COUNT + 1  # 38
f1, f2, f3 = skEnd + 1, skEnd + 2, skEnd + 3  # 42,43,44
label(f'A{f1}', '全項目3以上')
ws.cell(row=f1, column=2, value=f'=IF(COUNT(D{skStart}:D{skEnd})<{len(SKILL_ITEMS)},"",IF(MIN(D{skStart}:D{skEnd})>=3,"○","×"))')
label(f'A{f2}', '制作・現場で4以上の強み')
ws.cell(row=f2, column=2, value=f'=IF(COUNT(D{skStart}:D{skEnd})<{len(SKILL_ITEMS)},"",IF(MAX(D{fldStart}:D{skEnd})>=4,"○","×"))')
label(f'A{f3}', '1・2が残っている')
ws.cell(row=f3, column=2, value=f'=IF(COUNT(D{skStart}:D{skEnd})<{len(SKILL_ITEMS)},"",IF(COUNTIF(D{skStart}:D{skEnd},"<=2")>0,"有","無"))')

# 最終判定・処遇
pTitle = f3 + 2  # 46
section(pTitle, '■ 最終判定・処遇')
fin, raise_, now, new = pTitle + 1, pTitle + 2, pTitle + 3, pTitle + 4
label(f'A{fin}', '最終判定（粗利×バリュー）')
ws.cell(row=fin, column=2, value=f'=IF(OR($B$13="",$B${vj}=""),"",IFERROR(INDEX({MATRIX_VAL},MATCH($B$13,{MATRIX_ROW},0),MATCH($B${vj},{MATRIX_COL},0)),""))')
label(f'A{raise_}', '改定額')
ws.cell(row=raise_, column=2, value=f'=IF(OR($B${fin}="",$F$2=""),"",IFERROR(INDEX({RAISE_VAL},MATCH($B${fin},{RAISE_ROW},0),MATCH($F$2,{RAISE_COL},0)),""))')
ws.cell(row=raise_, column=2).number_format = '+#,##0;-#,##0;±0'
label(f'A{now}', '現在月給'); ws.cell(row=now, column=2, value='=$B$8').number_format = '#,##0'
label(f'A{new}', '改定後月給')
ws.cell(row=new, column=2, value=f'=IF(OR($B${raise_}="",$B${now}=""),"",$B${now}+$B${raise_})').number_format = '#,##0'

# 昇格・降格メモ
gTitle = new + 2  # 52
section(gTitle, '■ 昇格・降格メモ（連続条件は履歴を見て評価会議で判断）')
label(f'A{gTitle+1}', '昇格への本人意思'); inp(f'B{gTitle+1}')
label(f'A{gTitle+2}', '昇格後のアサイン・支援体制'); wide_input(gTitle + 2)
label(f'A{gTitle+3}', '降格の懸念・回復状況'); wide_input(gTitle + 3)

# 面談メモ
mTitle = gTitle + 5  # 57
section(mTitle, '■ 面談メモ（次の半期に向けて）')
memo = ['自分の粗利', '不足している要素', '次に取り組むこと', '伸ばすスキル', '必要な案件経験', '会社・チームに求めるサポート']
for i, mlabel in enumerate(memo):
    rr = mTitle + 1 + i
    label(f'A{rr}', mlabel); wide_input(rr); ws.row_dimensions[rr].height = 32

# 転記用
tTitle = mTitle + 1 + len(memo) + 1  # 65
section(tTitle, '■ 転記用（値をコピーして会議用シートへ貼り付け）')
t = tTitle + 1
label(f'A{t}', '粗利実績 →「粗利入力」F列'); ws.cell(row=t, column=2, value='=$B$11')
label(f'A{t+1}', 'バリュー自己 →「バリュー評価」C列')
for k, r in enumerate(range(vStart, vEnd + 1)):
    ws.cell(row=t + 1, column=2 + k, value=f'=$B${r}')
label(f'A{t+2}', 'バリュー長 →「バリュー評価」M列')
for k, r in enumerate(range(vStart, vEnd + 1)):
    ws.cell(row=t + 2, column=2 + k, value=f'=$C${r}')
label(f'A{t+3}', 'スキル自己 →「スキル評価」C列')
for k, r in enumerate(range(skStart, skEnd + 1)):
    ws.cell(row=t + 3, column=2 + k, value=f'=$B${r}')
label(f'A{t+4}', 'スキル長 →「スキル評価」M列')
for k, r in enumerate(range(skStart, skEnd + 1)):
    ws.cell(row=t + 4, column=2 + k, value=f'=$C${r}')
for rr in range(t, t + 5):
    ws.cell(row=rr, column=1).font = GREY9

# データ入力規則
dv03 = DataValidation(type='list', formula1='"0,1,2,3"', allow_blank=True); dv03.error = '0〜3で入力'
dv15 = DataValidation(type='list', formula1='"1,2,3,4,5"', allow_blank=True); dv15.error = '1〜5で入力'
dvg = DataValidation(type='list', formula1='"B,A1,A2,L1,L2,M1,M2"', allow_blank=True)
dvk = DataValidation(type='list', formula1='"正社員,パートナー"', allow_blank=True)
dvy = DataValidation(type='list', formula1='"1,2,3"', allow_blank=True)
dvi = DataValidation(type='list', formula1='"○,×,—"', allow_blank=True)
ws.add_data_validation(dv03); ws.add_data_validation(dv15)
ws.add_data_validation(dvg); ws.add_data_validation(dvk); ws.add_data_validation(dvy); ws.add_data_validation(dvi)
dv03.add(f'B{vStart}:C{vEnd}')
dv15.add(f'B{skStart}:C{skEnd}')
dvk.add('B5'); dvg.add('B6'); dvy.add('D6'); dvi.add(f'B{gTitle+1}')

# E:F 補助列を隠す
ws.column_dimensions['E'].hidden = True
ws.column_dimensions['F'].hidden = True
ws.freeze_panes = 'A2'

wb.active = wb['面談シート']
out = '/home/user/evaluation/templates/個人面談シート.xlsx'
import os
os.makedirs(os.path.dirname(out), exist_ok=True)
wb.save(out)
print('saved:', out)
